from __future__ import annotations

import sys
import os
import re
import zipfile
from xml.etree import ElementTree as ET
from pathlib import Path
from copy import deepcopy

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips


REFERENCE = Path(r"D:\qianrushi\智能乐谱辅助系统_本地视觉识别与演奏抓拍模块设计文档.docx")
OUTPUT = Path(r"D:\qianrushi\smartscore_s3_audio\智能乐谱辅助系统_本地音频采集与音乐识别子模块设计文档.docx")
WORK_DIR = Path(r"D:\qianrushi\smartscore_s3_audio\.docx_work")
FLOW_IMAGE = WORK_DIR / "music_recognition_flow.png"

BODY_WIDTH_TWIPS = 9070


def set_run_font(run, east_asia="Microsoft YaHei", latin="Calibri", size=None, bold=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def configure_style(style, *, east_asia, latin, size, bold, alignment,
                    before, after, line_spacing=None, first_indent_twips=0,
                    keep_with_next=False, keep_together=False):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    pf = style.paragraph_format
    pf.alignment = alignment
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Twips(first_indent_twips) if first_indent_twips else None
    if line_spacing is not None:
        pf.line_spacing = line_spacing
    pf.keep_with_next = keep_with_next
    pf.keep_together = keep_together


def configure_styles(doc):
    configure_style(
        doc.styles["Normal"], east_asia="Microsoft YaHei", latin="Calibri",
        size=10.5, bold=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        before=0, after=5, line_spacing=1.15, first_indent_twips=420,
    )
    configure_style(
        doc.styles["Heading 1"], east_asia="Microsoft YaHei", latin="Calibri",
        size=16, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT,
        before=16, after=8, keep_with_next=True, keep_together=True,
    )
    configure_style(
        doc.styles["Heading 2"], east_asia="Microsoft YaHei", latin="Calibri",
        size=13, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT,
        before=12, after=6, keep_with_next=True, keep_together=True,
    )
    configure_style(
        doc.styles["Heading 3"], east_asia="Microsoft YaHei", latin="Calibri",
        size=11, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT,
        before=8, after=4, keep_with_next=True, keep_together=True,
    )
    configure_style(
        doc.styles["Caption"], east_asia="Microsoft YaHei", latin="Calibri",
        size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER,
        before=4, after=4, line_spacing=1.0, keep_together=True,
    )
    if "Table Title" not in [style.name for style in doc.styles]:
        table_title = doc.styles.add_style("Table Title", WD_STYLE_TYPE.PARAGRAPH)
        table_title.base_style = doc.styles["Normal"]
    configure_style(
        doc.styles["Table Title"], east_asia="Microsoft YaHei", latin="Calibri",
        size=10.5, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER,
        before=8, after=4, line_spacing=1.0, keep_with_next=True,
        keep_together=True,
    )
    if "List Number" in [style.name for style in doc.styles]:
        list_style = doc.styles["List Number"]
        list_style.font.name = "Calibri"
        list_style.font.size = Pt(10.5)
        rpr = list_style.element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        list_style.paragraph_format.space_after = Pt(3)
        list_style.paragraph_format.line_spacing = 1.15


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def make_flow_image(path: Path):
    width, height = 1800, 360
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    font = ImageFont.truetype(str(font_path), 42)
    small = ImageFont.truetype(str(font_path), 32)

    labels = [
        ("双麦声音输入", "IM68A130 × 2"),
        ("音频数字化", "ES7210 / I2S"),
        ("双路预处理", "校准·高通·噪声门"),
        ("音频特征分析", "YIN + FFT"),
        ("分类与稳定判定", "单音·双音·和弦"),
    ]
    margin = 54
    arrow_gap = 54
    box_width = (width - 2 * margin - arrow_gap * 4) // 5
    box_height = 190
    y0 = 78
    palette = ["#EAF2F8", "#E8F5E9", "#FFF3E0", "#F3E5F5", "#E3F2FD"]
    border = "#315A7D"
    for index, (title, subtitle) in enumerate(labels):
        x0 = margin + index * (box_width + arrow_gap)
        x1 = x0 + box_width
        y1 = y0 + box_height
        draw.rounded_rectangle((x0, y0, x1, y1), radius=20, fill=palette[index],
                               outline=border, width=5)
        title_box = draw.textbbox((0, 0), title, font=font)
        subtitle_box = draw.textbbox((0, 0), subtitle, font=small)
        draw.text(((x0 + x1 - (title_box[2] - title_box[0])) / 2, y0 + 48),
                  title, font=font, fill="#17324D")
        draw.text(((x0 + x1 - (subtitle_box[2] - subtitle_box[0])) / 2, y0 + 116),
                  subtitle, font=small, fill="#455A64")
        if index < len(labels) - 1:
            start_x = x1 + 10
            end_x = x1 + arrow_gap - 10
            center_y = y0 + box_height // 2
            draw.line((start_x, center_y, end_x, center_y), fill=border, width=6)
            draw.polygon(
                [(end_x, center_y), (end_x - 22, center_y - 16),
                 (end_x - 22, center_y + 16)], fill=border
            )
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(300, 300))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="7A8793", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths, vertical_margin=80):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths[index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell, top=vertical_margin, bottom=vertical_margin)


def format_table(table, widths, narrative_columns=(), font_size=9.5,
                 vertical_margin=80):
    set_table_geometry(table, widths, vertical_margin=vertical_margin)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index == 0:
                set_cell_shading(cell, "315A7D")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F6F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = None
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if col_index in narrative_columns and row_index != 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    set_run_font(run, size=font_size, bold=(row_index == 0))
                    if row_index == 0:
                        run.font.color.rgb = None
                        color = OxmlElement("w:color")
                        color.set(qn("w:val"), "FFFFFF")
                        run._element.get_or_add_rPr().append(color)


def add_text(doc, text, style="Normal", *, keep_with_next=False):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    if keep_with_next:
        paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_table_title(doc, text):
    return add_text(doc, text, style="Table Title", keep_with_next=True)


def add_caption(doc, text):
    return add_text(doc, text, style="Caption", keep_with_next=False)


def add_result_table(doc):
    add_table_title(doc, "表 3.X-1  内部识别结果及含义")
    rows = [
        ("结果类型", "判定含义", "后续用途"),
        ("SINGLE", "稳定单音；给出频率、MIDI、音名、音分和置信度", "旋律跟随与单音评价"),
        ("INTERVAL", "两个稳定且相对独立的音级", "双音练习分析"),
        ("CHORD", "符合模板的常见大三或小三和弦", "和弦练习分析"),
        ("SILENCE", "两路输入均低于各自噪声门", "静音复位与无演奏判定"),
        ("UNKNOWN", "有声音活动，但削波、置信度不足或频谱不稳定", "抑制误识别并等待稳定结果"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    format_table(table, [1700, 4870, 2500], narrative_columns=(1, 2),
                 font_size=9.0, vertical_margin=45)


def add_hardware_table(doc):
    add_table_title(doc, "表 4.X-1  音频模块硬件接口与关键配置")
    rows = [
        ("对象/信号", "ESP32-S3连接或参数", "实现说明"),
        ("主控", "ESP32-S3", "双核 240 MHz；当前工程不启用 PSRAM"),
        ("模拟麦克风", "CN1/MIC1、CN2/MIC2", "两块 IM68A130 单端模拟输出，经板上交流耦合接入 ES7210"),
        ("ES7210 I2C", "SDA=GPIO4，SCL=GPIO5", "400 kHz；7 位地址 0x40；启动时校验 0x7210 芯片标识"),
        ("I2S 时钟", "MCLK=GPIO9，BCLK=GPIO10，WS=GPIO11", "24 kHz、16 位有效采样、双槽标准 I2S；MCLK=6.144 MHz"),
        ("I2S 数据", "DIN=GPIO12", "同步接收 MIC1 与 MIC2 两个时隙"),
        ("ES7210 INT", "GPIO42", "仅配置为输入，当前识别链路不启用中断"),
        ("输入增益", "30 dB 固定增益", "保留削波检测；运行期间不执行软件 AGC"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    format_table(table, [1900, 2700, 4470], narrative_columns=(2,))


def add_task_table(doc):
    add_table_title(doc, "表 4.X-2  识别模块任务与资源组织")
    rows = [
        ("执行单元", "核心/优先级/栈", "主要职责", "资源约束"),
        ("AudioCaptureTask", "Core 0 / 22 / 4096 B", "读取 I2S 双槽数据、拆分两路采样并投递采集块", "3 个固定采集块；队列满时记录丢块，不在采集路径做频谱运算"),
        ("MusicDspTask", "Core 1 / 12 / 6144 B", "校准、预处理、通道选择、YIN、FFT、分类与稳定投票", "使用静态环形缓冲和 FFT 工作区；循环内避免逐帧大块动态分配"),
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    format_table(table, [2050, 1800, 2500, 2720], narrative_columns=(2, 3), font_size=9.0)


def add_numbered_step(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.keep_with_next = False
    paragraph.add_run(text)
    return paragraph


def set_update_fields(doc):
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def ensure_first_page_header_footer(section):
    section.different_first_page_header_footer = True
    for source, target in (
        (section.header, section.first_page_header),
        (section.footer, section.first_page_footer),
    ):
        target_element = target._element
        for child in list(target_element):
            target_element.remove(child)
        for child in source._element:
            target_element.append(deepcopy(child))


def prune_unused_document_media(docx_path: Path):
    rels_path = "word/_rels/document.xml.rels"
    document_path = "word/document.xml"
    relationship_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(docx_path, "r") as source:
        document_xml = source.read(document_path)
        rels_root = ET.fromstring(source.read(rels_path))
        used_ids = {
            value.decode("utf-8")
            for value in re.findall(rb'r:(?:embed|link)="([^"]+)"', document_xml)
        }
        removed_targets = set()
        for relationship in list(rels_root):
            rel_type = relationship.attrib.get("Type", "")
            rel_id = relationship.attrib.get("Id", "")
            target = relationship.attrib.get("Target", "")
            if rel_type.endswith("/image") and rel_id not in used_ids:
                rels_root.remove(relationship)
                if target.startswith("media/"):
                    removed_targets.add("word/" + target)
        new_rels = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)
        temp_path = docx_path.with_suffix(".pruned.docx")
        with zipfile.ZipFile(temp_path, "w") as target_archive:
            for item in source.infolist():
                if item.filename in removed_targets:
                    continue
                data = new_rels if item.filename == rels_path else source.read(item.filename)
                target_archive.writestr(item, data)
    os.replace(temp_path, docx_path)


def main():
    if not REFERENCE.exists():
        raise SystemExit(f"reference not found: {REFERENCE}")
    make_flow_image(FLOW_IMAGE)
    doc = Document(REFERENCE)
    clear_body(doc)
    configure_styles(doc)
    set_update_fields(doc)
    ensure_first_page_header_footer(doc.sections[0])

    doc.core_properties.title = "智能乐谱辅助系统——本地音频采集与音乐识别子模块设计"
    doc.core_properties.subject = "智能乐谱辅助系统总设计文档子模块"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""

    add_text(doc, "3.X 本地音频采集与音乐识别功能设计", style="Heading 1")

    add_text(doc, "3.X.1 模块定位与主要功能", style="Heading 2")
    add_text(doc, "本模块是智能乐谱辅助系统的本地听觉感知子模块，面向电子琴扬声器播放和近场演奏场景，通过双麦克风采集声音，在 ESP32-S3 上完成音频预处理、基频估计、频谱分析和稳定分类。模块不依赖云端音频服务，也不上传原始 PCM 数据，可在演奏过程中持续生成音高与常见复音识别结果，为后续跟谱、评分和练习分析提供输入。")
    add_text(doc, "模块主要完成四项功能：同步采集两路麦克风信号并验证硬件状态；依据各通道噪声底和信号质量进行预处理与通道处理；联合 YIN 与 FFT 判断单音、双音和常见大/小三和弦；通过时间投票、起音检测和异常抑制输出稳定结果。")

    add_text(doc, "3.X.2 功能组成与数据流程", style="Heading 2")
    add_text(doc, "音频链路按照“声音输入—数字化采集—双路预处理—特征分析—分类稳定”的顺序运行。两块 IM68A130 分别连接 ES7210 的 MIC1、MIC2 模拟输入，ES7210 将两路信号转换为标准双槽 I2S 数据。采集层从同一帧中同步拆分两个时隙，避免把未对齐的采样直接进行时域叠加。")
    add_text(doc, "两路信号独立完成启动校准、去均值、高通滤波、噪声门和削波检查。单音基频分析使用质量较高的通道，复音频谱分析综合两路频谱证据；随后分类器依据周期性、谐波解释率、独立音级数量和模板置信度形成候选结果，并通过时间稳定机制输出最终类型。")
    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.first_line_indent = None
    picture_paragraph.paragraph_format.space_before = Pt(4)
    picture_paragraph.paragraph_format.space_after = Pt(0)
    picture_paragraph.paragraph_format.keep_with_next = True
    run = picture_paragraph.add_run()
    inline_shape = run.add_picture(str(FLOW_IMAGE), width=Inches(6.05))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", "双麦声音输入经 ES7210、双路预处理、YIN 与 FFT 分析后形成稳定音乐识别结果的流程图")
    add_caption(doc, "图 3.X-1  本地音频采集与音乐识别模块功能流程")

    add_text(doc, "3.X.3 单音、双音及和弦识别逻辑", style="Heading 2")
    add_text(doc, "单音识别以 YIN 周期检测为主。算法在时域中计算差分函数和累积均值归一化差分函数，搜索首个满足阈值的局部最小值，并通过抛物线插值提高周期估计精度。所得频率统一按 A4=440 Hz 转换为 MIDI、音名、八度和音分偏差。")
    add_text(doc, "为了兼顾电子琴音色和实际扬声器频响，单音候选采用多路径准入。周期置信度和谐波解释率同时达标时走严格路径；当频谱峰与 YIN 音高或音级一致时，可由频谱支持路径确认；在信号高于噪声门、未削波且音域合理的前提下，高置信度 YIN 还可作为旋律单音补充路径。该设计避免仅凭某一个频谱比例否决稳定的周期音，同时保留对噪声和错误八度的约束。")
    add_text(doc, "双音与和弦识别以 FFT 频谱为依据。系统先提取真实物理局部峰，再将峰值映射到 MIDI 音高并归属其 2～5 次谐波，避免把同一主瓣或单音谐波重复当作独立音符。两个占主导且能量相对均衡的独立音级形成 INTERVAL；至少三个稳定音级满足模板时，识别为常见大三或小三和弦。")

    add_text(doc, "3.X.4 稳定判定与异常处理", style="Heading 2")
    add_text(doc, "瞬时候选不会直接作为最终结果。分类器保存最近 5 次候选，单音、双音或和弦至少获得 3 票后才判为稳定；音量出现满足条件的快速上升时记录新的起音，用于区分同音连续弹奏。静音会清除历史状态，短时频谱泄漏或衰减阶段则允许已经稳定的结果在满足连续性条件时短暂保持。")
    add_text(doc, "当两路信号均低于各自噪声门时输出 SILENCE；输入削波、周期估计不可信、音级相互冲突或频谱持续不稳定时输出 UNKNOWN。双麦模式下，某一路信号过弱或削波时，质量比较会使另一通道承担主要的基频判断，避免单路异常直接中断识别。")
    add_result_table(doc)

    add_text(doc, "4.X 本地音频采集与音乐识别模块实现", style="Heading 1")

    add_text(doc, "4.X.1 硬件接口与音频采集", style="Heading 2")
    add_text(doc, "模块主控采用 ESP32-S3，两块 IM68A130 模拟麦克风分别通过板上交流耦合网络接入 ES7210 的 MIC1 与 MIC2。ES7210 使用 I2C 完成配置，使用 I2S0 向 ESP32-S3 输出两路同步采样。当前工程采样率为 24 kHz，每个时隙包含 16 位有效样本，共两个标准 I2S 时隙；固定输入增益为 30 dB，运行期间不启用软件自动增益。")
    add_text(doc, "启动时程序扫描 I2C 地址 0x40，读取芯片 ID 与版本寄存器，并校验芯片标识 0x7210。随后回读输入使能、增益、工作模式和串行数据格式，关闭不用于外部模拟麦克风供电的 MICBIAS12。双麦模式还会短暂隔离两个输入并比较时隙活动度；若硬件身份、关键寄存器回读或输入恢复失败，则不启动识别任务。")
    add_hardware_table(doc)

    add_text(doc, "4.X.2 双麦预处理与通道处理", style="Heading 2")
    add_text(doc, "AudioCaptureTask 每次读取 1024 个双槽采样帧，将同一 I2S 帧中的 MIC1、MIC2 数据分别写入固定采集块。DSP 端把 int16_t 样本转换为浮点数，逐帧去除均值，并为每路维护两级约 50 Hz 的一阶高通滤波状态，以减小直流漂移和低频机械扰动。")
    add_text(doc, "上电后约 1 s 内，两路分别估计 RMS 噪声底。实际噪声门取 max(0.0008, noise_rms×2.5)，并只在后续较安静的帧中向较低噪声水平缓慢恢复，避免把持续琴音吸收到噪声模型。每帧同时计算 RMS、峰值、削波率和直流偏置；峰值接近满量程或超过阈值的样本比例过高时，当前帧被标记为削波。")
    add_text(doc, "当前实现采用混合双麦策略：YIN 从质量评分较高的一路取得连续时域窗口；挑战通道的评分连续 3 帧高出当前通道至少 0.15 后才切换，以减少来回跳变。FFT 支路分别计算两路频谱，再依据两路 RMS 比例进行加权融合。这样既能为周期检测保留连续、清晰的单路波形，又能在复音判断时利用两处麦克风的频谱覆盖。")

    add_text(doc, "4.X.3 YIN、FFT及分类算法实现", style="Heading 2")
    add_text(doc, "YIN 使用 2048 点窗口，分析频率范围为 65～2000 Hz，累积均值归一化差分阈值为 0.20。程序按标准 YIN 规则选择首个达标局部最小值，不再偏好两倍周期，从而降低把 A4 等音高误判低一个八度的概率。周期结果的有效性、置信度和 MIDI 音域共同作为单音准入条件。")
    add_text(doc, "频谱分析使用 ESP-DSP 的 4096 点浮点 FFT 和 Hann 窗，hop 与采集块一致，均为 1024 点。程序只在 65～2000 Hz 有效频段内提取局部峰，通过抛物线插值得到更精确的峰值频率，并根据噪声底、相对得分和峰值突出度过滤弱候选。对每个候选基频继续检查 2～5 次谐波，以形成谐波解释率并清除重复音级。")
    add_text(doc, "分类器首先判断有效和弦或双音，再判断单音多路径准入。严格单音路径要求 YIN 置信度不低于 0.70 且谐波解释率不低于 0.74；频谱精确支持路径使用 0.88 的 YIN 置信度；旋律补充路径仅在 MIDI 48～84、信号高于噪声门且未削波时启用，其中强 YIN 阈值为 0.90，频谱同音级支持路径阈值为 0.80。有效复音与强单音发生冲突时，高置信度周期性和频谱证据共同决定是否由单音覆盖。")
    add_text(doc, "稳定投票窗口长度为 5，至少 3 次相同类型和身份才输出单音、双音或和弦。起音检测要求当前 RMS 相对上一帧至少上升 1.45 倍，且绝对增量不低于 0.0006；相邻起音之间至少间隔 180 ms。以上阈值均集中在 music_detector_config.h，便于依据目标硬件日志统一调整。")

    heading_44 = add_text(doc, "4.X.4 FreeRTOS任务与资源组织", style="Heading 2")
    heading_44.paragraph_format.page_break_before = True
    add_text(doc, "实时软件将高优先级采集与计算量较大的 DSP 分离。AudioCaptureTask 固定在 Core 0，只负责 I2S 接收、双槽拆分和采集块投递；MusicDspTask 固定在 Core 1，完成预处理、YIN、FFT、分类与诊断统计。采集队列使用三个固定块，队列暂时耗尽时记录丢块而不在中断或采集路径等待复杂处理。")
    add_text(doc, "工程当前不依赖 PSRAM。双路 PCM 帧、4096 点环形历史、YIN 工作窗口和 FFT 工作区均采用静态或启动阶段分配，DSP 循环不执行逐帧大块 malloc/free。诊断计数器记录 I2S 读取错误、队列深度、缓冲耗尽、DSP 平均/最大耗时、截止期丢失和任务栈余量，为后续实机调参提供依据。")
    add_task_table(doc)

    add_text(doc, "4.X.5 构建验证、实机测试方法与能力边界", style="Heading 2")
    add_text(doc, "当前工程面向 ESP32-S3，依赖 esp_codec_dev 与 esp-dsp，已生成可烧录应用镜像。软件构建能够验证组件依赖、内存链接和算法代码的完整性；项目同时提供单音准入门控测试用例，用于覆盖严格路径、高置信度 YIN、频谱支持以及低于噪声门、削波和越界音高拒绝等条件。构建与主机侧逻辑用例不能替代真实麦克风、扬声器、房间反射和演奏动态下的板端验证。")
    add_text(doc, "建议在目标 PCB 上按以下顺序验收：")
    add_numbered_step(doc, "静音上电并完成双路校准，检查 ES7210 身份、时隙映射、增益回读、两路噪声底和噪声门。")
    add_numbered_step(doc, "分别测试 C4、E4、G4、A4、C5 等单音，核对 MIDI、音名、八度与音分偏差，并连续演奏简单旋律。")
    add_numbered_step(doc, "测试 C4+B4 等双音及 C-E-G、D-F-A、G-B-D 等常见大/小三和弦，检查音级集合和稳定投票。")
    add_numbered_step(doc, "以拍手、人声、背景音乐、弱音和削波输入检查 UNKNOWN 与 SILENCE，统计静音误触和错误音符。")
    add_numbered_step(doc, "分别遮挡或断开 CN1、CN2，确认另一麦克风仍可承担主要分析，并检查通道切换是否稳定。")
    add_text(doc, "验收应记录正确音符率、静音误触次数、八度错误次数、起音延迟、双麦切换次数和 DSP 截止期丢失。当前模块适合初学者单音旋律、双音和常见三和弦识别；双音的音级通常比具体八度更可靠，三和弦以常见大/小模板为主。四音以上复杂复音、延音踏板造成的长时间重叠、强混响、多声源分离及未经统计的实机准确率不属于当前版本承诺。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    prune_unused_document_media(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
