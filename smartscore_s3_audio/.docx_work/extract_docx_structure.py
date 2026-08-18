from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def paragraph_record(paragraph: Paragraph) -> dict:
    fmt = paragraph.paragraph_format
    return {
        "kind": "paragraph",
        "style": paragraph.style.name if paragraph.style else "",
        "text": paragraph.text,
        "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
        "page_break_before": bool(fmt.page_break_before),
        "keep_with_next": bool(fmt.keep_with_next),
        "runs": [
            {
                "text": run.text,
                "bold": run.bold,
                "italic": run.italic,
                "font": run.font.name,
                "size_pt": run.font.size.pt if run.font.size else None,
            }
            for run in paragraph.runs
            if run.text
        ],
    }


def table_record(table: Table) -> dict:
    return {
        "kind": "table",
        "style": table.style.name if table.style else "",
        "rows": [
            ["\n".join(p.text for p in cell.paragraphs).strip() for cell in row.cells]
            for row in table.rows
        ],
    }


def header_footer_text(section) -> dict:
    return {
        "header": [p.text for p in section.header.paragraphs if p.text],
        "footer": [p.text for p in section.footer.paragraphs if p.text],
        "first_page_header": [p.text for p in section.first_page_header.paragraphs if p.text],
        "first_page_footer": [p.text for p in section.first_page_footer.paragraphs if p.text],
    }


def extract(path: Path, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document(path)
    blocks = []
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            blocks.append(paragraph_record(block))
        else:
            blocks.append(table_record(block))

    sections = []
    for idx, section in enumerate(doc.sections, start=1):
        section_data = {
            "index": idx,
            "width_twips": int(section.page_width.twips),
            "height_twips": int(section.page_height.twips),
            "top_margin_twips": int(section.top_margin.twips),
            "bottom_margin_twips": int(section.bottom_margin.twips),
            "left_margin_twips": int(section.left_margin.twips),
            "right_margin_twips": int(section.right_margin.twips),
            "header_distance_twips": int(section.header_distance.twips),
            "footer_distance_twips": int(section.footer_distance.twips),
            "different_first_page": bool(section.different_first_page_header_footer),
        }
        section_data.update(header_footer_text(section))
        sections.append(section_data)

    styles = []
    for style in doc.styles:
        if style.type != 1:
            continue
        pf = style.paragraph_format
        rfonts = style.element.rPr.rFonts if style.element.rPr is not None else None
        styles.append(
            {
                "name": style.name,
                "font": style.font.name,
                "font_ascii": rfonts.get(qn("w:ascii")) if rfonts is not None else None,
                "font_hansi": rfonts.get(qn("w:hAnsi")) if rfonts is not None else None,
                "font_east_asia": rfonts.get(qn("w:eastAsia")) if rfonts is not None else None,
                "font_cs": rfonts.get(qn("w:cs")) if rfonts is not None else None,
                "size_pt": style.font.size.pt if style.font.size else None,
                "bold": style.font.bold,
                "alignment": int(pf.alignment) if pf.alignment is not None else None,
                "space_before_pt": pf.space_before.pt if pf.space_before else None,
                "space_after_pt": pf.space_after.pt if pf.space_after else None,
                "line_spacing": float(pf.line_spacing) if isinstance(pf.line_spacing, (int, float)) else None,
                "first_line_indent_twips": int(pf.first_line_indent.twips) if pf.first_line_indent else None,
            }
        )

    result = {
        "path": str(path),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "section_count": len(doc.sections),
        "sections": sections,
        "styles": styles,
        "blocks": blocks,
    }
    (out_dir / "structure.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    text_lines = []
    for block in blocks:
        if block["kind"] == "paragraph":
            if block["text"].strip():
                text_lines.append(f"[{block['style']}] {block['text']}")
        else:
            text_lines.append(f"[TABLE style={block['style']}]")
            for row in block["rows"]:
                text_lines.append(" | ".join(row))
    (out_dir / "content.txt").write_text("\n".join(text_lines), encoding="utf-8")

    media_dir = out_dir / "media"
    media_dir.mkdir(exist_ok=True)
    package_inventory = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith("/"):
                continue
            data = archive.read(name)
            package_inventory.append(
                {
                    "path": name,
                    "size": len(data),
                    "sha256": hashlib.sha256(data).hexdigest(),
                }
            )
            if name.startswith("word/media/") and not name.endswith("/"):
                (media_dir / Path(name).name).write_bytes(data)
    (out_dir / "package_inventory.json").write_text(
        json.dumps(package_inventory, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: extract_docx_structure.py INPUT.docx OUTPUT_DIR")
    extract(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
